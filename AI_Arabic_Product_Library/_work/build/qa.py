#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Automated QA for the Arabic AI Product Library."""

import os, re, sys, json, unicodedata
import pymupdf

AR = re.compile("[؀-ۿ]")

# Arabic Presentation Forms — must NOT appear in extracted text.
PRESENTATION = re.compile("[ﭐ-﷿ﹰ-﻿]")

def norm(s):
    s = s.replace("\u200f", "").replace("\u200e", "")
    s = unicodedata.normalize("NFKC", s)
    return re.sub(r"\s+", " ", s).strip()

def pdf_report(pdf_path, canonical_path, png_dir=None, dpi=110):
    doc = pymupdf.open(pdf_path)
    r = {"pdf": os.path.basename(pdf_path), "pages": len(doc), "issues": [], "checks": {}}

    # ---- fonts + embedding ----
    fonts, not_embedded = set(), set()
    for p in doc:
        for f in p.get_fonts(full=True):
            fonts.add(f[3])
            if not f[1]:
                not_embedded.add(f[3])
    r["checks"]["fonts"] = sorted(fonts)
    r["checks"]["fonts_not_embedded"] = sorted(not_embedded)
    if not_embedded:
        r["issues"].append(("CRITICAL", "fonts not embedded: %s" % sorted(not_embedded)))

    # ---- per page: text present, no overflow, no presentation forms ----
    W, H = doc[0].rect.width, doc[0].rect.height
    empty, overflow, presforms = [], [], []
    MARGIN_L, MARGIN_R = 28.0, 28.0     # pt; page margin is 16mm = 45pt
    for i, p in enumerate(doc, 1):
        txt = p.get_text().strip()
        if not txt and not p.get_images() and not p.get_drawings():
            empty.append(i)
        if PRESENTATION.search(txt):
            presforms.append(i)
        for b in p.get_text("blocks"):
            x0, y0, x1, y1 = b[:4]
            if x0 < MARGIN_L - 2 or x1 > W - MARGIN_R + 2 or y1 > H - 12 or y0 < 8:
                # cover page is full-bleed by design
                if i != 1:
                    overflow.append((i, round(x0, 1), round(x1, 1), round(y0, 1), round(y1, 1)))
    r["checks"]["empty_pages"] = empty
    r["checks"]["overflow_blocks"] = overflow[:40]
    r["checks"]["presentation_form_pages"] = presforms
    if empty:
        r["issues"].append(("CRITICAL", "empty pages: %s" % empty))
    if overflow:
        r["issues"].append(("CRITICAL", "%d text blocks outside the safe area (first: %s)"
                            % (len(overflow), overflow[0])))
    if presforms:
        r["issues"].append(("CRITICAL",
                            "Arabic presentation forms in the text layer on pages %s "
                            "(copying would not yield real Arabic letters)" % presforms))

    # ---- selectable text on every content page ----
    no_text = [i for i, p in enumerate(doc, 1) if i > 1 and len(p.get_text().strip()) < 20]
    r["checks"]["pages_without_text"] = no_text
    if no_text:
        r["issues"].append(("MAJOR", "pages with almost no selectable text: %s" % no_text))

    # ---- links ----
    internal = external = 0
    tg_links, broken = [], []
    for i, p in enumerate(doc, 1):
        for l in p.get_links():
            if l["kind"] in (pymupdf.LINK_GOTO, pymupdf.LINK_NAMED):
                internal += 1
                tgt = l.get("page", -1)
                if not (0 <= tgt < len(doc)):
                    broken.append((i, l.get("kind"), tgt))
            elif l["kind"] == pymupdf.LINK_URI:
                external += 1
                if "t.me/PromptsArabic" in l.get("uri", ""):
                    tg_links.append((i, l["uri"]))
    r["checks"]["internal_links"] = internal
    r["checks"]["external_links"] = external
    r["checks"]["telegram_link_annotations"] = tg_links
    r["checks"]["broken_internal_links"] = broken
    if broken:
        r["issues"].append(("MAJOR", "broken internal links: %s" % broken[:5]))
    if not tg_links:
        r["issues"].append(("CRITICAL", "no clickable https://t.me/PromptsArabic link annotation"))

    # ---- copyright presence ----
    all_text = "\n".join(p.get_text() for p in doc).replace("\u200f", "").replace("\u200e", "")
    r["checks"]["owner_mentions"] = all_text.count("برومبتات عربية")
    r["checks"]["tg_text_mentions"] = all_text.count("t.me/PromptsArabic")
    if r["checks"]["owner_mentions"] == 0:
        r["issues"].append(("CRITICAL", "owner name missing from the PDF text"))
    if r["checks"]["tg_text_mentions"] == 0:
        r["issues"].append(("CRITICAL", "t.me/PromptsArabic missing from the PDF text"))

    # ---- search test (visual + text agreement) ----
    probes = ["برومبتات عربية", "t.me/PromptsArabic"]
    found = {}
    for t in probes:
        found[t] = sum(len(p.search_for(t)) for p in doc)
    r["checks"]["search_hits"] = found
    for t, n in found.items():
        if n == 0:
            r["issues"].append(("CRITICAL", "search failed for %r" % t))

    # ---- copy / search test against canonical source ----
    if canonical_path and os.path.exists(canonical_path):
        canon = open(canonical_path, encoding="utf-8").read()
        canon_words = sorted({w for w in re.findall(r"[ء-ي\u0640\u064B-\u0652]{4,}", canon)})
        haystack = norm(all_text)
        missing = [w for w in canon_words if w not in haystack]
        r["checks"]["canonical_arabic_words"] = len(canon_words)
        r["checks"]["extract_missing_words_count"] = len(missing)
        r["checks"]["extract_missing_words"] = missing[:25]
        if missing:
            r["issues"].append(("CRITICAL",
                "%d canonical Arabic words absent from the extracted text: %s"
                % (len(missing), missing[:10])))
        # reader-style find, on a spread sample across the whole document
        step = max(1, len(canon_words) // 120)
        sample = canon_words[::step]
        unfindable = [w for w in sample
                      if not any(p.search_for(w) for p in doc)]
        r["checks"]["search_sample_size"] = len(sample)
        r["checks"]["search_unfindable"] = unfindable[:25]
        if unfindable:
            r["issues"].append(("CRITICAL",
                "%d/%d sampled Arabic words are not findable by PDF search: %s"
                % (len(unfindable), len(sample), unfindable[:10])))
        # authoring rules that keep variables copy-safe
        digit_vars = sorted(set(re.findall(r"\[[^\]\n]*[0-9][^\]\n]*\]", canon)))
        r["checks"]["variables_with_digits"] = digit_vars
        if digit_vars:
            r["issues"].append(("MAJOR",
                "variables contain Western digits, which break bracket order when copied: %s"
                % digit_vars[:8]))

        # variables must survive copy with their brackets
        variables = sorted(set(re.findall(r"\[[ء-ي\u0640\u064B-\u0652 ]{2,30}\]", canon)))
        bad = [v for v in variables if v not in haystack or not any(p.search_for(v) for p in doc)]
        r["checks"]["variables"] = variables
        r["checks"]["variables_broken"] = bad
        if bad:
            r["issues"].append(("CRITICAL", "variables not intact in the PDF: %s" % bad))

    # ---- render every page ----
    if png_dir:
        os.makedirs(png_dir, exist_ok=True)
        for i, p in enumerate(doc, 1):
            p.get_pixmap(dpi=dpi).save(os.path.join(png_dir, "p%03d.png" % i))
        r["checks"]["pages_rendered"] = len(doc)
    doc.close()
    return r


def contact_sheet(png_dir, out_path, cols=5, thumb_w=300):
    from PIL import Image
    files = sorted(f for f in os.listdir(png_dir) if f.endswith(".png") and f.startswith("p"))
    if not files:
        return None
    ims = []
    for f in files:
        im = Image.open(os.path.join(png_dir, f))
        h = int(im.height * thumb_w / im.width)
        ims.append(im.resize((thumb_w, h), Image.LANCZOS))
    tw, th = thumb_w, max(i.height for i in ims)
    rows = (len(ims) + cols - 1) // cols
    pad = 8
    sheet = Image.new("RGB", (cols * (tw + pad) + pad, rows * (th + pad) + pad), "#DCD6CC")
    for idx, im in enumerate(ims):
        x = pad + (idx % cols) * (tw + pad)
        y = pad + (idx // cols) * (th + pad)
        sheet.paste(im, (x, y))
    sheet.save(out_path)
    return out_path


def docx_report(path):
    from docx import Document
    from docx.oxml.ns import qn
    d = Document(path)
    r = {"docx": os.path.basename(path), "issues": [], "checks": {}}
    paras = d.paragraphs
    r["checks"]["paragraphs"] = len(paras)
    r["checks"]["tables"] = len(d.tables)
    styles = sorted({p.style.name for p in paras if p.style})
    r["checks"]["styles_used"] = styles
    bidi = sum(1 for p in paras if p._p.find(qn("w:pPr")) is not None
               and p._p.pPr.find(qn("w:bidi")) is not None)
    r["checks"]["paragraphs_with_bidi"] = bidi
    sect_bidi = all(s._sectPr.find(qn("w:bidi")) is not None for s in d.sections)
    r["checks"]["sections_rtl"] = sect_bidi
    if not sect_bidi:
        r["issues"].append(("CRITICAL", "document section is not RTL"))
    txt = "\n".join(p.text for p in paras)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                txt += "\n" + c.text
    r["checks"]["owner_mentions"] = txt.count("برومبتات عربية")
    r["checks"]["tg_mentions"] = txt.count("t.me/PromptsArabic")
    if r["checks"]["owner_mentions"] == 0:
        r["issues"].append(("CRITICAL", "owner name missing from the DOCX"))
    if r["checks"]["tg_mentions"] == 0:
        r["issues"].append(("CRITICAL", "t.me/PromptsArabic missing from the DOCX"))
    if not any(s.startswith("APL") for s in styles):
        r["issues"].append(("MAJOR", "APL paragraph styles are not being applied"))
    return r


if __name__ == "__main__":
    print(json.dumps(pdf_report(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None),
                     ensure_ascii=False, indent=2))
